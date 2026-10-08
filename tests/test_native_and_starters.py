import json
import subprocess
from pathlib import Path

from docx import Document
from fastapi import Request
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Accomplishment, ReportTemplate
from app.services import native_picker
from app.services.reports import generate_docx, regenerate_docx
from app.services.starter_report_templates import STARTERS, seed_starter_templates
from test_auth_and_capture import token


def test_native_picker_auth_csrf_and_remote_restrictions(logged_in, monkeypatch):
    assert (
        logged_in.post("/files/pick-native", data={"csrf": "bad", "kind": "odt"}).status_code == 403
    )
    assert (
        logged_in.post(
            "/files/pick-native", data={"csrf": token(logged_in), "kind": "odt"}
        ).status_code
        == 403
    )
    for host, headers in [("192.0.2.1", []), ("127.0.0.1", [(b"x-forwarded-for", b"192.0.2.1")])]:
        request = Request(
            {
                "type": "http",
                "scheme": "http",
                "server": ("localhost", 8000),
                "path": "/",
                "headers": headers,
                "client": (host, 50000),
            }
        )
        assert not native_picker.available(request)
    monkeypatch.setattr(native_picker, "available", lambda request: True)
    monkeypatch.setattr(native_picker, "choose", lambda *args: {"canceled": True})
    response = logged_in.post("/files/pick-native", data={"csrf": token(logged_in), "kind": "odt"})
    assert response.json() == {"canceled": True}


def test_native_selection_is_revalidated(tmp_path, monkeypatch):
    root = tmp_path / "repository"
    root.mkdir()
    child = root / "evidence"
    child.mkdir()
    chosen = {"selected": str(child)}
    monkeypatch.setattr(
        native_picker.subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(args, 0, json.dumps(chosen), ""),
    )
    assert native_picker.choose("subdirectory", "", str(root)) == {"selected": "evidence"}
    chosen["selected"] = str(tmp_path)
    import pytest

    with pytest.raises(ValueError, match="inside"):
        native_picker.choose("subdirectory", "", str(root))
    chosen["selected"] = "//server/share/source.odt"
    with pytest.raises(ValueError, match="local drive"):
        native_picker.choose("odt", "", "")
    chosen.clear()
    chosen["canceled"] = True
    assert native_picker.choose("odt", "", "") == {"canceled": True}


def test_eight_starters_preserve_edits_and_generate_distinct_layouts():
    with SessionLocal() as session:
        assert seed_starter_templates(session.connection()) == 8
        session.commit()
        templates = list(session.scalars(select(ReportTemplate)))
        assert len(templates) == len(STARTERS) == 8
        templates[0].content = "My customized notes"
        templates[0].name = "My renamed starter"
        session.commit()
        assert seed_starter_templates(session.connection()) == 0
        assert session.get(ReportTemplate, templates[0].id).content == "My customized notes"
        record = Accomplishment(
            title="Evidence",
            action="Reviewed settings",
            metric="24 settings",
            impact="Supported exception decisions",
            supporting_narrative="Documented rationale",
        )
        session.add(record)
        session.commit()
        for layout in ("detailed", "impact_first", "compact"):
            report = generate_docx(session, "Test layout", [record], layout=layout)
            document = Document(report.output_path)
            text = "\n".join(p.text for p in document.paragraphs)
            assert "Documented rationale" in text
            if layout == "impact_first":
                assert text.index("Impact") < text.index("Action")
            elif layout == "compact":
                assert "Action: Reviewed settings" in text
            else:
                assert text.index("Action") < text.index("Impact")
            from app.models import ReportItem

            items = list(
                session.scalars(select(ReportItem).where(ReportItem.report_id == report.id))
            )
            regenerate_docx(session, report, items)
            assert "\n".join(p.text for p in Document(report.output_path).paragraphs) == text
            assert Path(report.output_path).is_file()
