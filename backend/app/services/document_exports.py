"""Local document downloads. Separate from the restricted Markdown/Git pipeline."""

from datetime import date, datetime, time
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from odf import style, table, text
from odf.opendocument import OpenDocumentText
from sqlalchemy.orm import Session
from xlsxwriter import Workbook

from app.models import Accomplishment
from app.services import accomplishments, reporting_year

MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "odt": "application/vnd.oasis.opendocument.text",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def candidates(
    session: Session,
    scope: str = "current",
    project_id: str = "",
    include_archived: bool = False,
    include_unfinished: bool = False,
    query: str = "",
) -> list[Accomplishment]:
    settings = reporting_year.preferences(session)
    period = reporting_year.current_period(settings)
    records = accomplishments.search(
        session, query=query, include_archived=include_archived, project_id=project_id
    )
    if not include_unfinished:
        records = [record for record in records if record.status in {"completed", "archived"}]
    if scope == "current":
        records = reporting_year.current_records(records, period)
    elif scope != "all":
        try:
            start = date.fromisoformat(scope)
            selected = reporting_year.period_for(start, settings)
        except (ValueError, OverflowError) as exc:
            raise ValueError("Choose a valid reporting year.") from exc
        if start != selected.start:
            raise ValueError("Choose a valid reporting year.")
        records = [
            record
            for record in records
            if record.date_completed and selected.start <= record.date_completed <= selected.end
        ]
    return sorted(
        records, key=lambda record: (record.date_completed or date.min, record.title), reverse=True
    )


def record_metadata(record: Accomplishment) -> str:
    return (
        f"Start: {record.date_started or 'Not set'}   Completed: {record.date_completed or 'Not set'}\n"
        f"Projects: {', '.join(project.name for project in record.projects) or 'None'}\n"
        f"Status: {record.status}   Sensitivity: {record.sensitivity}"
    )


def word_document(records: list[Accomplishment], title: str, include_notes: bool) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ["Normal", "Title", "Heading 1"]:
        font = document.styles[name].font
        font.name, font.color.rgb = "Calibri", RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10)
    document.add_paragraph(title, "Title")
    document.add_paragraph(f"{len(records)} accomplishments")
    for record in records:
        document.add_heading(record.title, level=1)
        document.add_paragraph(record_metadata(record))
        grid = document.add_table(rows=1, cols=3)
        grid.autofit = False
        for col in grid.columns:
            col.width = Inches(9.7 / 3)
        header = grid.rows[0]
        repeat = OxmlElement("w:tblHeader")
        header._tr.get_or_add_trPr().append(repeat)
        for cell, label in zip(header.cells, ["Action", "Metric", "Impact"], strict=True):
            cell.text = label
            cell.paragraphs[0].runs[0].bold = True
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EDF2")
            cell._tc.get_or_add_tcPr().append(shade)
        for cell, value in zip(
            grid.add_row().cells, [record.action, record.metric, record.impact], strict=True
        ):
            cell.text = value or "Not provided"
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{edge}")
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        grid._tbl.tblPr.append(borders)
        if include_notes:
            for label, value in [
                ("Supporting notes", record.supporting_narrative),
                ("Original note", record.raw_note),
            ]:
                if value:
                    document.add_paragraph(f"{label}: {value}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def odt_document(records: list[Accomplishment], title: str, include_notes: bool) -> bytes:
    document = OpenDocumentText()
    layout = style.PageLayout(name="Landscape")
    layout.addElement(
        style.PageLayoutProperties(
            pagewidth="11in", pageheight="8.5in", printorientation="landscape", margin="0.65in"
        )
    )
    document.automaticstyles.addElement(layout)
    document.masterstyles.addElement(style.MasterPage(name="Standard", pagelayoutname=layout))
    body = style.Style(name="Body", family="paragraph")
    body.addElement(style.TextProperties(fontfamily="Calibri", fontsize="10pt", color="#000000"))
    body.addElement(style.ParagraphProperties(marginbottom="0.08in"))
    document.styles.addElement(body)
    title_style = style.Style(name="Title", family="paragraph", parentstylename=body)
    title_style.addElement(style.TextProperties(fontsize="22pt", fontweight="bold"))
    document.styles.addElement(title_style)
    heading = style.Style(name="Record", family="paragraph", parentstylename=body)
    heading.addElement(style.TextProperties(fontsize="14pt", fontweight="bold"))
    heading.addElement(style.ParagraphProperties(margintop="0.15in", keepwithnext="always"))
    document.styles.addElement(heading)
    cell_style = style.Style(name="Cell", family="table-cell")
    cell_style.addElement(style.TableCellProperties(border="0.5pt solid #D9D9D9", padding="0.08in"))
    document.automaticstyles.addElement(cell_style)
    header_style = style.Style(name="Header", family="table-cell", parentstylename=cell_style)
    header_style.addElement(style.TableCellProperties(backgroundcolor="#E8EDF2"))
    header_style.addElement(style.TextProperties(fontweight="bold"))
    document.automaticstyles.addElement(header_style)
    column = style.Style(name="Column", family="table-column")
    column.addElement(style.TableColumnProperties(columnwidth="3.233in"))
    document.automaticstyles.addElement(column)
    document.text.addElement(text.P(stylename=title_style, text=title))
    document.text.addElement(text.P(stylename=body, text=f"{len(records)} accomplishments"))
    for index, record in enumerate(records):
        document.text.addElement(text.H(outlinelevel=1, stylename=heading, text=record.title))
        for line in record_metadata(record).splitlines():
            document.text.addElement(text.P(stylename=body, text=line))
        grid = table.Table(name=f"Accomplishment{index + 1}")
        grid.addElement(table.TableColumn(stylename=column, numbercolumnsrepeated=3))
        headers = table.TableHeaderRows()
        row = table.TableRow()
        for label in ["Action", "Metric", "Impact"]:
            cell = table.TableCell(stylename=header_style)
            cell.addElement(text.P(stylename=body, text=label))
            row.addElement(cell)
        headers.addElement(row)
        grid.addElement(headers)
        row = table.TableRow()
        for value in [record.action, record.metric, record.impact]:
            cell = table.TableCell(stylename=cell_style)
            for line in (value or "Not provided").splitlines():
                cell.addElement(text.P(stylename=body, text=line))
            row.addElement(cell)
        grid.addElement(row)
        document.text.addElement(grid)
        if include_notes:
            for label, value in [
                ("Supporting notes", record.supporting_narrative),
                ("Original note", record.raw_note),
            ]:
                if value:
                    for line in f"{label}: {value}".splitlines():
                        document.text.addElement(text.P(stylename=body, text=line))
    output = BytesIO()
    document.write(output)
    return output.getvalue()


def excel_document(records: list[Accomplishment], title: str, include_notes: bool) -> bytes:
    output = BytesIO()
    with Workbook(
        output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    ) as workbook:
        sheet = workbook.add_worksheet("Accomplishments")
        sheet.hide_gridlines(2)
        title_format = workbook.add_format({"bold": True, "font_size": 18})
        header = workbook.add_format(
            {"bold": True, "bg_color": "#243C53", "font_color": "#FFFFFF", "text_wrap": True}
        )
        body = workbook.add_format({"text_wrap": True, "valign": "top"})
        dates = workbook.add_format({"num_format": "mm/dd/yyyy", "valign": "top"})
        columns = [
            "Title",
            "Start date",
            "Completion date",
            "Projects",
            "Action",
            "Metric",
            "Impact",
            "Status",
            "Sensitivity",
        ]
        if include_notes:
            columns += ["Supporting notes", "Original note"]
        sheet.merge_range(0, 0, 0, 3, title, title_format)
        sheet.write(1, 0, f"{len(records)} accomplishments")
        sheet.write_row(3, 0, columns, header)
        sheet.set_row(3, 32)
        sheet.set_column(0, 0, 34)
        sheet.set_column(1, 2, 15)
        sheet.set_column(3, 3, 26)
        sheet.set_column(4, 6, 48)
        sheet.set_column(7, 8, 22)
        if include_notes:
            sheet.set_column(9, 10, 55)
        for index, record in enumerate(records, 4):
            values = [
                record.title,
                record.date_started,
                record.date_completed,
                ", ".join(project.name for project in record.projects),
                record.action,
                record.metric,
                record.impact,
                record.status,
                record.sensitivity,
            ]
            if include_notes:
                values += [record.supporting_narrative, record.raw_note]
            for col, value in enumerate(values):
                if isinstance(value, date):
                    sheet.write_datetime(index, col, datetime.combine(value, time()), dates)
                else:
                    # Explicit strings prevent formula and external-link injection.
                    result = sheet.write_string(index, col, str(value or ""), body)
                    if result == -2:
                        raise ValueError(
                            "A field exceeds Excel's cell length limit. Use Word or ODT instead."
                        )
            sheet.set_row(index, 90)
        sheet.autofilter(3, 0, 3 + len(records), len(columns) - 1)
        sheet.freeze_panes(4, 1)
    return output.getvalue()


def generate(records: list[Accomplishment], title: str, format: str, include_notes: bool) -> bytes:
    return {"docx": word_document, "odt": odt_document, "xlsx": excel_document}[format](
        records, title, include_notes
    )
