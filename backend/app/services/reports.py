from __future__ import annotations

from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Inches, Pt
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Accomplishment, AuditEvent, Report, ReportItem


def _add_record(document: WordDocument, snapshot: dict[str, object]) -> None:
    document.add_heading(str(snapshot.get("title") or "Untitled accomplishment"), level=1)
    for heading, field in (("Action", "action"), ("Metric", "metric"), ("Impact", "impact")):
        document.add_heading(heading, level=2)
        document.add_paragraph(str(snapshot.get(field) or "[Not provided]"))
    supporting_narrative = str(snapshot.get("supporting_narrative") or "")
    if supporting_narrative:
        document.add_heading("Supporting details", level=2)
        document.add_paragraph(supporting_narrative)


def _build_document(
    title: str, report_type: str, snapshots: list[dict[str, object]], template_content: str
) -> WordDocument:
    document = Document()
    document.sections[0].top_margin = Inches(0.75)
    style = document.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(10)
    document.add_heading(title, 0)
    document.add_paragraph(f"CareerForge {report_type.title()} report")
    if template_content:
        document.add_heading("Report notes", level=1)
        document.add_paragraph(template_content)
    for snapshot in snapshots:
        _add_record(document, snapshot)
    return document


def _safe_name(title: str) -> str:
    return "-".join("".join(char if char.isalnum() else " " for char in title).split()).lower() or (
        "careerforge-report"
    )


def generate_docx(
    session: Session,
    title: str,
    records: list[Accomplishment],
    report_type: str = "custom",
    template_id: str = "",
    template_content: str = "",
) -> Report:
    report = Report(
        title=title,
        report_type=report_type,
        filters={"template_id": template_id, "template_content": template_content},
    )
    session.add(report)
    session.flush()
    snapshots: list[dict[str, object]] = []
    for position, record in enumerate(records, start=1):
        snapshot: dict[str, object] = {
            "title": record.title,
            "action": record.action,
            "metric": record.metric,
            "impact": record.impact,
            "supporting_narrative": record.supporting_narrative,
        }
        snapshots.append(snapshot)
        session.add(
            ReportItem(
                report_id=report.id,
                accomplishment_id=record.id,
                position=position,
                source_snapshot=snapshot,
            )
        )
    output = get_settings().data_dir / "reports" / f"{_safe_name(title)}-{report.id[:8]}.docx"
    _build_document(title, report_type, snapshots, template_content).save(str(output))
    report.output_path = str(output)
    session.add(
        AuditEvent(
            event_type="report.generated",
            metadata_json={"report_id": report.id, "record_count": len(records)},
        )
    )
    session.commit()
    return report


def regenerate_docx(session: Session, report: Report, items: list[ReportItem]) -> None:
    """Rebuild a report from its immutable item snapshots after its order changes."""
    filters = report.filters if isinstance(report.filters, dict) else {}
    template_content = str(filters.get("template_content") or "")
    snapshots = [item.source_snapshot for item in items]
    _build_document(report.title, report.report_type, snapshots, template_content).save(report.output_path)
    session.add(
        AuditEvent(
            event_type="report.regenerated",
            metadata_json={"report_id": report.id, "record_count": len(items)},
        )
    )
    session.commit()
